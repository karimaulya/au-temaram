"""
xlsx (chat log) -> src/story/chats.generated.json
Jalankan: python scripts/convert.py Temaram.xlsx src/story/chats.generated.json
Butuh: pip install openpyxl
"""
import json, re, sys, datetime
from openpyxl import load_workbook

SRC, OUT = sys.argv[1], sys.argv[2]
SELF = "Asahi"

# nama sheet -> chatId (sheet "X (2)" otomatis jadi chunk ke-2 dari chat "X")
CHAT_IDS = {
    "GRUP Ujang Kedu": "group",
    "Jae berisik": "jae",
    "Owner Penginapan": "owner",
    "Bang Ucuk": "ucuk",
    "Jeo": "jeo",
    "Mia": "mia",
}

def fmt_time(v):
    if isinstance(v, (datetime.time, datetime.datetime)):
        return v.strftime("%H:%M")
    if isinstance(v, float):  # fraksi hari ala Excel
        m = round(v * 24 * 60)
        return f"{m // 60:02d}:{m % 60:02d}"
    return str(v) if v is not None else ""

wb = load_workbook(SRC, read_only=True, data_only=False)
chats, chunks = {}, []
counter = {}

for ws in wb:
    base = re.sub(r"\s*\(\d+\)$", "", ws.title)
    chat_id = CHAT_IDS.get(base, re.sub(r"\W+", "-", base.lower()))
    counter[chat_id] = counter.get(chat_id, 0) + 1
    chunk_id = f"{chat_id}-{counter[chat_id]}"
    chats.setdefault(chat_id, {"id": chat_id, "title": base, "isGroup": chat_id == "group"})

    msgs = []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if not r or r[0] is None:
            continue
        _id, sender, text, t = r[0], r[1], r[2], r[3]
        cue = r[5] if len(r) > 5 and r[5] else None
        text = str(text).strip()
        kind = "text"
        if text == "Shared location": kind = "location"
        elif text.lower().startswith("file "): kind = "file"
        elif re.match(r"^https?://", text): kind = "link"
        m = {
            # id di sheet bentrok antar-sheet (msg-15, msg-38), jadi diberi prefix chunk
            "id": f"{chunk_id}:{_id}",
            "sender": sender,
            # isSelf di sheet 2+ berupa rumus; hitung sendiri dari nama pengirim
            "isSelf": sender == SELF,
            "text": text,
            "kind": kind,
        }
        t_str = fmt_time(t)
        if re.match(r"^\d{2}:\d{2}$", t_str): m["time"] = t_str
        else: m["dateLabel"] = t_str
        if cue:
            m["cue"] = re.sub(r"^\((.*)\)$", r"\1", str(cue).strip())
        msgs.append(m)

    chunks.append({"id": chunk_id, "chatId": chat_id, "sheet": ws.title, "messages": msgs})

json.dump({"chats": list(chats.values()), "chunks": chunks}, open(OUT, "w", encoding="utf-8"),
          ensure_ascii=False, indent=2)
print(f"{len(chunks)} chunk, {sum(len(c['messages']) for c in chunks)} pesan -> {OUT}")
